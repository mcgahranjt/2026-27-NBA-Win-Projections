"""
Google Sheet sync

Copies the script-owned tabs from NBA_Spread_Calculator.xlsx into the Google Sheet as plain values,
the same as copying and using paste-values. Run it after the other scripts have updated the workbook.

What gets written:
  Power Rankings      A1:J31 only (the Spread Calculator looks up rows 2-31; column K onward is untouched)
  Team Ratings        A1:L31 (team table plus league ORtg/pace and notes in K-L; Matchup Predictor reads these)
  Player Projections  whole tab replaced
  Player Ranges       whole tab replaced (created if missing)
  Game Projections    whole tab replaced from output_power/game_projections.csv (created if missing)
  Model Grades        whole tab replaced, once nba_grade.py has written it (created if missing)
Every other tab is left alone. Power Rankings and Team Ratings are only written if their header row
in the Google Sheet still matches, so a moved or renamed column stops the sync instead of breaking
the calculator.

One-time setup: a Google Cloud service account with the Sheets API enabled, its JSON key saved here
as google_service_account.json, the Google Sheet shared with the service account's email as an
Editor, and pip3 install gspread.

Usage:
  python3 sync_to_sheets.py            # update the Google Sheet
  python3 sync_to_sheets.py --dry-run  # show what would be written, without connecting
"""

import datetime as dt
import os
import sys

import pandas as pd
from openpyxl import load_workbook

import nba_power_ratings as PR

SHEET_ID = "1jdNhA_p575HAXA4alnl0pilNyRLIfhS941fEEWVFJKs"
KEY_PATH = "google_service_account.json"
GAME_PROJECTIONS = os.path.join(PR.OUT_DIR, "game_projections.csv")

# (Google tab, workbook sheet, fixed range or None for the whole tab, header columns to check)
FIXED = [("Power Rankings", "Power Rankings", "A1:J31", 10),
         ("Team Ratings", "Team Ratings", "A1:L31", 9)]
WHOLE = [("Player Projections", "Player Projections"),
         ("Player Ranges", "Player Ranges"),
         ("Model Grades", "Model Grades")]
FILTERED = {"Player Projections", "Player Ranges", "Game Projections"}
GAME_HEADERS = {"date": "Date", "away": "Away", "home": "Home", "proj_margin": "Home margin",
                "proj_total": "Total", "home_win_prob": "Home win prob", "as_of": "Projected on"}


def clean(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    return v.item() if hasattr(v, "item") else v


def sheet_values(ws, max_row=None, max_col=None):
    rows = ws.iter_rows(min_row=1, max_row=max_row or ws.max_row, max_col=max_col or ws.max_column,
                        values_only=True)
    return [[clean(v) for v in r] for r in rows]


def range_size(a1):
    from openpyxl.utils.cell import range_boundaries
    c1, r1, c2, r2 = range_boundaries(a1)
    return r2 - r1 + 1, c2 - c1 + 1


def gather():
    """Everything to send, read from the workbook and output files: [(tab, range or None, values, header_cols)]."""
    if not os.path.exists(PR.CALCULATOR_PATH):
        raise SystemExit(f"Can't find {PR.CALCULATOR_PATH}.")
    wb = load_workbook(PR.CALCULATOR_PATH, read_only=True)
    jobs = []
    for tab, src, rng, check in FIXED:
        if src in wb.sheetnames:
            rows, cols = range_size(rng)
            jobs.append((tab, rng, sheet_values(wb[src], rows, cols), check))
    for tab, src in WHOLE:
        if src in wb.sheetnames:
            vals = [r for r in sheet_values(wb[src]) if any(v != "" for v in r)]
            if vals:
                jobs.append((tab, None, vals, 0))
    if os.path.exists(GAME_PROJECTIONS):
        gp = pd.read_csv(GAME_PROJECTIONS, encoding="utf-8-sig", dtype={"game_id": str})
        gp = gp.drop(columns=["game_id"]).rename(columns=GAME_HEADERS)[list(GAME_HEADERS.values())]
        jobs.append(("Game Projections", None,
                     [list(gp.columns)] + [[clean(v) for v in r] for r in gp.itertuples(index=False)], 0))
    wb.close()
    return jobs


def connect():
    if not os.path.exists(KEY_PATH):
        raise SystemExit(f"Can't find {KEY_PATH}. See the setup steps at the top of this script.")
    try:
        import gspread
    except ImportError:
        raise SystemExit("gspread isn't installed. Run: pip3 install gspread")
    return gspread.service_account(filename=KEY_PATH).open_by_key(SHEET_ID)


def write(book, tab, rng, values, check):
    titles = [w.title for w in book.worksheets()]
    rows, cols = len(values), max(len(r) for r in values)
    if rng:
        if tab not in titles:
            print(f"  {tab}: tab not found in the Google Sheet, skipped.")
            return
        ws = book.worksheet(tab)
        header = (ws.row_values(1) + [""] * check)[:check]
        if [str(h) for h in header] != [str(v) for v in values[0][:check]]:
            print(f"  {tab}: header row doesn't match the workbook, so it wasn't updated.\n"
                  f"     Google: {header}\n     Workbook: {values[0][:check]}")
            return
        ws.update(range_name=rng, values=values, value_input_option="RAW")
        print(f"  {tab}: {rng} updated.")
        return
    if tab in titles:
        ws = book.worksheet(tab)
        ws.clear()
        if ws.row_count < rows or ws.col_count < cols:
            ws.resize(rows=max(ws.row_count, rows), cols=max(ws.col_count, cols))
    else:
        ws = book.add_worksheet(title=tab, rows=max(rows, 100), cols=max(cols, 26))
        print(f"  {tab}: new tab created.")
    ws.update(range_name="A1", values=values, value_input_option="RAW")
    ws.freeze(rows=1)
    if tab in FILTERED:
        ws.set_basic_filter(f"A1:{chr(64 + min(cols, 26))}{rows}")
    print(f"  {tab}: {rows - 1} rows written.")


def main():
    jobs = gather()
    if "--dry-run" in sys.argv:
        print("Dry run - nothing sent. Would write:")
        for tab, rng, values, check in jobs:
            where = rng or f"whole tab, {len(values)} rows x {max(len(r) for r in values)} columns"
            print(f"  {tab}: {where}; first row {values[0][:6]}")
        return
    book = connect()
    print(f"Updating Google Sheet '{book.title}' ({dt.datetime.now():%Y-%m-%d %H:%M}):")
    for job in jobs:
        try:
            write(book, *job)
        except Exception as e:
            print(f"  {job[0]}: failed ({e.__class__.__name__}: {str(e)[:120]})")


if __name__ == "__main__":
    main()
