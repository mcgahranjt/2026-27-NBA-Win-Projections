"""
NBA model grading

Checks the projections against what actually happened, so you can see whether the model is any
good and which settings need adjusting.

Player props (from nba_player_props.py's history files vs. box scores):
  - MAE: average miss per stat, next to a simple baseline (the player's season average so far).
    The model should beat the baseline; if it doesn't for a stat, that stat needs work.
  - Bias: average projection minus actual. Positive = projecting too high.
  - Range check: how often the result fell under the median (target ~50%), below the 10th
    percentile and above the 90th (each target ~10%; lower for low-count stats like steals and
    blocks, where many results land exactly on the range edges). If results fall outside the
    range much more often than that, the ranges are too narrow; much less often, too wide.
  - Players projected who didn't play are counted separately (props are voided for them).

Games (from nba_power_ratings.py's game_projection_history.csv vs. final scores), using the latest
projection made on or before game day:
  - Margin MAE vs. a baseline of "home team by 2.3", straight-up winner accuracy, Brier score for
    win probabilities (lower is better; 0.25 is a coin flip), total MAE and bias, and the share of
    games that went over the model's projected total (about 50% if the totals don't lean either way).

Preseason anchor test (--tune): replays this season day by day. Each day's games are predicted
using only earlier games, once for each anchor size, to see how much the preseason ratings should
count. Results are reported, not applied - change PRIOR_GAMES / PACE_PRIOR_GAMES in
nba_power_ratings.py yourself if a different value keeps winning over several weeks.

Outputs (output_grades/): player_grades.csv, player_grades_daily.csv, game_grades.csv,
game_results.csv, anchor_tuning.csv (with --tune). If NBA_Spread_Calculator.xlsx is in this folder,
the Model Grades sheet is replaced with the summary. Close the workbook first.

Usage:
  python3 nba_grade.py          # grade everything so far
  python3 nba_grade.py --tune   # also run the preseason anchor test (best after ~4 weeks of games)
"""

import os
import sys

import numpy as np
import pandas as pd
from scipy.stats import norm

import nba_player_props as P
import nba_power_ratings as PR
import nba_season_sim as S

OUT_DIR = "output_grades"
STAT_ORDER = ["MIN", "PTS", "REB", "AST", "3PM", "STL", "BLK", "TOV", "PRA", "P+R", "P+A", "R+A"]
LOG_COLS = {"PTS": "PTS", "REB": "REB", "AST": "AST", "3PM": "FG3M", "STL": "STL", "BLK": "BLK", "TOV": "TOV"}
LOW_COUNT = {"3PM", "STL", "BLK", "TOV"}
BASELINE_MIN_GAMES = 3
ANCHORS = [1, 5, 10, 20, 30, 40]   # 1 = almost no preseason anchor
PACE_ANCHORS = [1, 5, 10, 20, 30]
TUNE_SKIP_DAYS = 7        # the first week has too few earlier games to predict from


# ------------------------------------------------------------ player props
def actual_player_stats(log):
    a = log[["PLAYER_ID", "GAME_DATE", "MIN"] + P.STATS].copy()
    a = a.rename(columns={v: k for k, v in LOG_COLS.items()})
    a["PRA"], a["P+R"] = a.PTS + a.REB + a.AST, a.PTS + a.REB
    a["P+A"], a["R+A"] = a.PTS + a.AST, a.REB + a.AST
    a = a.sort_values(["PLAYER_ID", "GAME_DATE"])
    # baseline: his average in this season's earlier games
    for s in STAT_ORDER:
        prior = a.groupby("PLAYER_ID")[s].transform(lambda x: x.shift().expanding().mean())
        n = a.groupby("PLAYER_ID").cumcount()
        a[f"base_{s}"] = prior.where(n >= BASELINE_MIN_GAMES)
    a["date"] = a.GAME_DATE.dt.date.astype(str)
    return a


def grade_players(log):
    hist_path = os.path.join(P.OUT_DIR, "projection_history.csv")
    if not os.path.exists(hist_path):
        print("No player projections saved yet (run nba_player_props.py on game days).")
        return None, None
    proj = pd.read_csv(hist_path, encoding="utf-8-sig")
    actual = actual_player_stats(log)
    played_dates = set(actual.date)
    proj = proj[proj.date.isin(played_dates)]
    if proj.empty:
        print("No completed games yet for the saved player projections.")
        return None, None
    m = proj.merge(actual, on=["date", "PLAYER_ID"], how="left", suffixes=("_proj", ""))
    dnp = m.MIN.isna()
    print(f"Player projections graded: {(~dnp).sum()} player-games ({dnp.sum()} projected players didn't play).")
    m = m[~dnp]

    rh_path = os.path.join(P.OUT_DIR, "range_history.csv")
    ranges = pd.read_csv(rh_path, encoding="utf-8-sig") if os.path.exists(rh_path) else None

    rows, daily = [], []
    for s in STAT_ORDER:
        err = m[f"{s}_proj"] - m[s]
        has_base = m[f"base_{s}"].notna()
        row = {"stat": s, "n": len(m), "MAE": err.abs().mean(), "bias": err.mean(),
               "baseline_n": int(has_base.sum()),
               "MAE_same_games": err[has_base].abs().mean(),
               "baseline_MAE": (m[f"base_{s}"] - m[s])[has_base].abs().mean()}
        if ranges is not None and s != "MIN":
            r = ranges[ranges.Stat == s].merge(m[["date", "PLAYER_ID", s]], on=["date", "PLAYER_ID"])
            if len(r):
                row["pct_under_median"] = 100 * ((r[s] < r.Median).mean() + 0.5 * (r[s] == r.Median).mean())
                row["pct_below_P10"] = 100 * (r[s] < r.P10).mean()
                row["pct_above_P90"] = 100 * (r[s] > r.P90).mean()
        rows.append(row)
        d = m.assign(abs_err=err.abs()).groupby("date").agg(n=("abs_err", "size"), MAE=("abs_err", "mean"))
        daily.append(d.reset_index().assign(stat=s))
    g = pd.DataFrame(rows)
    g["model_beats_baseline"] = np.where(g.baseline_n == 0, "",
                                         np.where(g.MAE_same_games < g.baseline_MAE, "yes", "no"))
    g["range_verdict"] = g.apply(range_verdict, axis=1)
    return g.round(2), pd.concat(daily, ignore_index=True).round(2)


def range_verdict(r):
    if pd.isna(r.get("pct_below_P10", np.nan)) or r.n < 100:
        return "not enough games yet" if r.stat != "MIN" else ""
    outside = r.pct_below_P10 + r.pct_above_P90
    if outside > 28:
        return "ranges too narrow"
    if outside < (6 if r.stat in LOW_COUNT else 12):
        return "ranges too wide"
    if abs(r.pct_under_median - 50) > 7:
        return "median too " + ("high" if r.pct_under_median > 50 else "low")
    return "ok"


# ------------------------------------------------------------------ games
def grade_games(sched):
    path = os.path.join(PR.OUT_DIR, "game_projection_history.csv")
    if not os.path.exists(path):
        print("No game projections saved yet (nba_power_ratings.py saves them on each run).")
        return None, None
    h = pd.read_csv(path, encoding="utf-8-sig", dtype={"game_id": str})
    h = h[h.as_of <= h.date].sort_values("as_of").groupby("game_id").tail(1)
    res = sched[sched.home_pts.notna()][["game_id", "home_pts", "away_pts"]]
    m = h.merge(res, on="game_id")
    if m.empty:
        print("No completed games yet for the saved game projections.")
        return None, None
    m["margin"] = m.home_pts - m.away_pts
    m["total"] = m.home_pts + m.away_pts
    m["home_won"] = (m.margin > 0).astype(int)
    m["margin_err"] = m.proj_margin - m.margin
    m["total_err"] = m.proj_total - m.total
    g = pd.DataFrame([{
        "games": len(m),
        "margin_MAE": m.margin_err.abs().mean(),
        "baseline_margin_MAE": (PR.HOME_COURT - m.margin).abs().mean(),
        "winner_pct": 100 * ((m.proj_margin > 0) == (m.margin > 0)).mean(),
        "brier": ((m.home_win_prob - m.home_won) ** 2).mean(),
        "total_MAE": m.total_err.abs().mean(),
        "total_bias": m.total_err.mean(),
        # share of games whose actual total went over the model's projected total (about 50% if unbiased)
        "pct_over_total": 100 * (m.total > m.proj_total).mean(),
    }]).round(3)
    print(f"Games graded: {len(m)}.")
    return g, m


# ------------------------------------------------------------ anchor test
def anchor_test(teams, sched):
    played = sched[sched.home_pts.notna()].copy()
    played["date"] = pd.to_datetime(played.date)
    if len(played) < 150:
        print(f"\nAnchor test: only {len(played)} games played; results will be noisy until ~200+.")
    teams, lg_o0, lg_p0 = PR.add_priors(teams, PR.fetch_last_season_advanced())
    played["game_id"] = played.game_id.astype(str).str.zfill(10)
    played = played.merge(PR.fetch_box_log(), on="game_id", how="left")
    has_poss = played.poss.notna()
    for c in ["home_pts", "away_pts", "rest_adj"]:
        played[c] = played[c].astype(float)
    dates = sorted(played.date.unique())
    dates = [d for d in dates if d >= dates[0] + pd.Timedelta(days=TUNE_SKIP_DAYS)]

    orig = PR.PRIOR_GAMES, PR.PACE_PRIOR_GAMES
    rows = []
    try:
        for kind, grid in [("efficiency", ANCHORS), ("pace", PACE_ANCHORS)]:
            for v in grid:
                PR.PRIOR_GAMES = v if kind == "efficiency" else orig[0]
                PR.PACE_PRIOR_GAMES = v if kind == "pace" else orig[1]
                errs = []
                for d in dates:
                    before = played[played.date < d]
                    bp = before[has_poss[before.index]]
                    lg_o, lg_p = lg_o0, lg_p0
                    if len(bp):
                        w = len(bp) / (len(bp) + 60)
                        lg_o = w * 100 * (bp.home_pts.sum() + bp.away_pts.sum()) / (2 * bp.poss.sum()) + (1 - w) * lg_o0
                        lg_p = w * bp.poss.mean() + (1 - w) * lg_p0
                    g = before.assign(poss=before.poss.fillna(lg_p))
                    t = teams.copy()
                    t["off"], t["dfn"] = PR.fit_off_def(t, g, lg_o, d)
                    t["pace"] = PR.fit_pace(t, g[has_poss[g.index]], lg_p, d)
                    r = t.set_index("TEAM_ID")
                    for x in played[played.date == d].itertuples():
                        mm = PR.matchup(r.loc[x.home_id], r.loc[x.away_id], lg_o, lg_p, x.rest_adj)
                        errs.append((mm["margin"] - (x.home_pts - x.away_pts), mm["total"] - (x.home_pts + x.away_pts)))
                e = np.array(errs)
                current = v == (orig[0] if kind == "efficiency" else orig[1])
                rows.append({"anchor_for": kind, "anchor_games": v, "games_predicted": len(e),
                             "margin_MAE": np.abs(e[:, 0]).mean(), "total_MAE": np.abs(e[:, 1]).mean(),
                             "current_setting": "yes" if current else ""})
    finally:
        PR.PRIOR_GAMES, PR.PACE_PRIOR_GAMES = orig
    out = pd.DataFrame(rows).round(3)
    for kind, col, name in [("efficiency", "margin_MAE", "PRIOR_GAMES"), ("pace", "total_MAE", "PACE_PRIOR_GAMES")]:
        k = out[out.anchor_for == kind]
        best, cur = k.loc[k[col].idxmin()], k[k.current_setting == "yes"].iloc[0]
        gain = cur[col] - best[col]
        msg = "keep the current setting" if gain < 0.05 else f"consider {name} = {int(best.anchor_games)} if it holds up"
        print(f"  {kind} anchor: best {int(best.anchor_games)} games ({col} {best[col]:.2f}) vs current "
              f"{int(cur.anchor_games)} ({cur[col]:.2f}) -> {msg}")
    return out


# ------------------------------------------------------------------ main
def write_workbook(pg, gg):
    if not os.path.exists(PR.CALCULATOR_PATH) or (pg is None and gg is None):
        return
    from openpyxl import load_workbook
    try:
        wb = load_workbook(PR.CALCULATOR_PATH)
    except PermissionError:
        print(f"Couldn't open {PR.CALCULATOR_PATH} - close it in Excel and run again.")
        return
    if "Model Grades" in wb.sheetnames:
        ws = wb["Model Grades"]
        ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet("Model Grades")
    for title, df in [("Games", gg), ("Player projections", pg)]:
        if df is None:
            continue
        ws.append([title])
        ws.append(list(df.columns))
        for row in df.itertuples(index=False):
            ws.append([None if pd.isna(v) else (v.item() if hasattr(v, "item") else v) for v in row])
        ws.append([])
    ws.column_dimensions["A"].width = 16
    try:
        wb.save(PR.CALCULATOR_PATH)
        print(f"Updated {PR.CALCULATOR_PATH} (Model Grades).")
    except PermissionError:
        print(f"Couldn't save {PR.CALCULATOR_PATH} - close it in Excel and run again.")


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    teams = PR.use_weighted_preseason(S.load_teams())
    sched = S.fetch_schedule(set(teams.TEAM_ID))
    if sched is None:
        raise SystemExit("Couldn't get the schedule.")
    sched["game_id"] = sched.game_id.astype(str).str.zfill(10)
    sched = PR.back_to_back_adjustment(sched)

    gg, games = grade_games(sched)
    if gg is not None:
        gg.to_csv(os.path.join(OUT_DIR, "game_grades.csv"), index=False)
        games.to_csv(os.path.join(OUT_DIR, "game_results.csv"), index=False, encoding="utf-8-sig")
        print("\nGames\n" + gg.to_string(index=False))

    pg, daily = grade_players(P.season_game_log())
    if pg is not None:
        pg.to_csv(os.path.join(OUT_DIR, "player_grades.csv"), index=False)
        daily.to_csv(os.path.join(OUT_DIR, "player_grades_daily.csv"), index=False)
        print("\nPlayer projections (MAE_same_games and baseline_MAE use only games where the player "
              f"had {BASELINE_MIN_GAMES}+ earlier games)\n" + pg.to_string(index=False))

    if "--tune" in sys.argv:
        print("\nPreseason anchor test (each day predicted from earlier games only):")
        at = anchor_test(teams, sched)
        at.to_csv(os.path.join(OUT_DIR, "anchor_tuning.csv"), index=False)
        print(at.to_string(index=False))

    write_workbook(pg, gg)
    print(f"\nSaved results to ./{OUT_DIR}")


if __name__ == "__main__":
    main()
