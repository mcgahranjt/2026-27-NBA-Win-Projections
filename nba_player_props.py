"""
NBA daily player stat projections

Projects every likely player's minutes, points, rebounds, assists, threes, steals, blocks and
turnovers for a day's games, plus the combos (PRA, P+R, P+A, R+A) and double-double odds.
Built to help spot player prop angles; it doesn't compare against sportsbook lines yet.

How it works:
  1. Per-minute rates. Each player's rate for every stat starts from his last two seasons
     (last season counts double), pulled toward the league average when he played few minutes.
     This season's games, weighted toward recent ones, take over as they come in.
  2. Minutes. A preseason expectation (Version B's projected minutes, or last season's) blended
     with this season's recent minutes, so role changes show up within a week or two. Each team's
     active players are then fit to 240 minutes, so when someone is out his minutes go to teammates.
  3. Who's playing. Out and Doubtful players come from ESPN's public injury feed (which follows the
     teams' official reports). data/injury_overrides.csv lets you mark anyone Out or Active by hand.
  4. Game context, reused from the team ratings (nothing is re-estimated here):
       - pace: more expected possessions = more of every stat
       - each team's projected score from the Matchup Projection math; players' points, threes and
         assists are scaled so the team adds up to it. That brings in the opponent's defense, and
         when a good player is out his teammates absorb most of his scoring.
       - out players lower their team's projected score using Version B's player ratings
       - lopsided games: starters lose some minutes when the projected margin is over 10
  5. Ranges. Each player's game is simulated 4,000 times (minutes vary, then each stat varies
     around its expected value), which gives medians, 10th-90th percentile ranges and
     double-double odds. Projections assume the player plays (props are voided if he doesn't).

Run nba_power_ratings.py first on the same day so the team ratings are current.

Outputs (output_props/):
  player_projections_<date>.csv  - one row per player
  player_ranges_<date>.csv       - mean, 10th/25th/50th/75th/90th percentiles for every stat
  projection_history.csv         - every day's projections, for grading against results later
  range_history.csv              - every day's ranges (nba_grade.py checks their width)
If NBA_Spread_Calculator.xlsx is in this folder, the Player Projections and Player Ranges sheets
are replaced with the day's numbers (no other sheet is touched). Close the workbook first.

Usage:
  python3 nba_player_props.py              # today's games
  python3 nba_player_props.py 2026-10-22   # another date
"""

import datetime as dt
import os
import sys

import numpy as np
import pandas as pd
import requests

import nba_power_ratings as PR
import nba_season_sim as S
import nba_wins_version_b as V
from nba_wins_version_b import name_key

OUT_DIR = "output_props"
STATS = ["PTS", "REB", "AST", "FG3M", "STL", "BLK", "TOV"]
LABELS = {"PTS": "PTS", "REB": "REB", "AST": "AST", "FG3M": "3PM", "STL": "STL", "BLK": "BLK", "TOV": "TOV"}
COMBOS = {"PRA": ["PTS", "REB", "AST"], "P+R": ["PTS", "REB"], "P+A": ["PTS", "AST"], "R+A": ["REB", "AST"]}
SCORING = ["PTS", "FG3M", "AST"]    # scaled to the team's projected score
VOLUME = ["REB", "STL", "BLK", "TOV"]  # scaled by pace only

# Rates: how many minutes of league-average play a player's history is blended with
HIST_PRIOR_MIN = {"PTS": 250, "REB": 250, "AST": 250, "FG3M": 400, "STL": 600, "BLK": 500, "TOV": 400}
# ...and how many minutes his history counts for once this season's games come in
CUR_PRIOR_MIN = {"PTS": 400, "REB": 400, "AST": 400, "FG3M": 600, "STL": 800, "BLK": 700, "TOV": 600}
LAST_SEASON_WEIGHT, TWO_SEASONS_WEIGHT = 1.0, 0.5
ROOKIE_RATE = 0.85         # players with no NBA history start at 85% of league-average rates
RATE_HALF_LIFE = 45        # days
# Minutes
MIN_PRIOR_GAMES = 5        # the preseason minutes expectation counts as 5 games
MIN_HALF_LIFE = 15         # days; minutes react to role changes faster than rates
PART_PRIOR_GAMES = 3       # same idea for how often a player gets into games
UNKNOWN_MPG = 10.0
TEAM_MINUTES = 240
MAX_MIN = 42               # hard cap on anyone's projected minutes
REDIST_CAP = 38            # extra minutes from injuries can't push a player past 38 (or 1 over his norm)
MIN_TO_SHOW = 6            # players projected under 6 minutes are left out of the output
DEPTH_WEIGHT = 0.5         # when a team has more minutes than 240, half the cut comes off the end of the
                           # bench (Version B's depth chart) and half evenly; best on the start of 2025-26
BLOWOUT_START, BLOWOUT_PER_PT = 10, 0.008   # starters lose 0.8% of minutes per point of margin past 10
STARTER_MIN = 28
REPLACEMENT_BPM = -2.0     # the player who replaces an injured one is about a -2 (Version B's prior)
# Simulation
N_SIMS = 4000
SEED = 7
DISPERSION = {"PTS": 20, "REB": 15, "AST": 12, "FG3M": 8, "STL": 10, "BLK": 4, "TOV": 15}  # higher = steadier
OUT_STATUSES = ("out", "doubtful", "suspension", "suspended", "inactive")

ESPN_INJURIES = "https://site.api.espn.com/apis/site/v2/sports/basketball/nba/injuries"
ABBR = {1610612737: "ATL", 1610612738: "BOS", 1610612739: "CLE", 1610612740: "NOP", 1610612741: "CHI",
        1610612742: "DAL", 1610612743: "DEN", 1610612744: "GSW", 1610612745: "HOU", 1610612746: "LAC",
        1610612747: "LAL", 1610612748: "MIA", 1610612749: "MIL", 1610612750: "MIN", 1610612751: "BKN",
        1610612752: "NYK", 1610612753: "ORL", 1610612754: "IND", 1610612755: "PHI", 1610612756: "PHX",
        1610612757: "POR", 1610612758: "SAC", 1610612759: "SAS", 1610612760: "OKC", 1610612761: "TOR",
        1610612762: "UTA", 1610612763: "MEM", 1610612764: "WAS", 1610612765: "DET", 1610612766: "CHA"}


# ------------------------------------------------------------------ data
def season_label(start_year):
    return f"{start_year}-{str(start_year + 1)[-2:]}"


def prev_season_totals(start_year):
    """A completed season's player totals from NBA.com (cached; past seasons don't change)."""
    season = season_label(start_year)
    path = os.path.join(S.DATA_DIR, f"player_totals_{season}.csv")
    if os.path.exists(path):
        return pd.read_csv(path)
    try:
        from nba_api.stats.endpoints import leaguedashplayerstats
        df = leaguedashplayerstats.LeagueDashPlayerStats(season=season, per_mode_detailed="Totals",
                                                         timeout=60).get_data_frames()[0]
        df = df[["PLAYER_ID", "PLAYER_NAME", "TEAM_ID", "GP", "MIN"] + STATS]
        df.to_csv(path, index=False)
        return df
    except Exception as e:
        print(f"Couldn't get {season} player totals ({e.__class__.__name__}).")
        return pd.DataFrame(columns=["PLAYER_ID", "PLAYER_NAME", "TEAM_ID", "GP", "MIN"] + STATS)


def current_players():
    """Every player on a roster right now, with his team (refreshed each run, saved as a fallback)."""
    path = os.path.join(S.DATA_DIR, "current_players.csv")
    try:
        from nba_api.stats.endpoints import commonallplayers
        df = commonallplayers.CommonAllPlayers(is_only_current_season=1, season=S.SEASON,
                                               timeout=60).get_data_frames()[0]
        df = df[df.TEAM_ID != 0][["PERSON_ID", "DISPLAY_FIRST_LAST", "TEAM_ID"]]
        df.columns = ["PLAYER_ID", "PLAYER_NAME", "TEAM_ID"]
        df.to_csv(path, index=False)
        return df
    except Exception as e:
        if os.path.exists(path):
            print(f"Couldn't refresh rosters ({e.__class__.__name__}); using the saved copy.")
            return pd.read_csv(path)
        raise SystemExit("Couldn't get current rosters from NBA.com - try again in a few minutes.")


def season_game_log():
    """This season's player box scores so far."""
    path = os.path.join(S.DATA_DIR, f"player_game_log_{S.SEASON}.csv")
    cols = ["PLAYER_ID", "PLAYER_NAME", "TEAM_ID", "GAME_ID", "GAME_DATE", "MIN"] + STATS
    try:
        from nba_api.stats.endpoints import leaguegamelog
        log = leaguegamelog.LeagueGameLog(season=S.SEASON, player_or_team_abbreviation="P",
                                          season_type_all_star="Regular Season", timeout=60).get_data_frames()[0]
        log = log[cols] if not log.empty else pd.DataFrame(columns=cols)
        log.to_csv(path, index=False)
    except Exception as e:
        if not os.path.exists(path):
            print(f"Couldn't get this season's box scores ({e.__class__.__name__}); using history only.")
            return pd.DataFrame(columns=cols)
        print(f"Couldn't refresh box scores ({e.__class__.__name__}); using the saved copy.")
        log = pd.read_csv(path)
    log["GAME_DATE"] = pd.to_datetime(log.GAME_DATE)
    log["MIN"] = pd.to_numeric(log.MIN, errors="coerce").fillna(0)
    return log[log.MIN > 0]


def injury_report():
    """Injury statuses by player name key. ESPN's feed, then the manual overrides file on top."""
    rows = []
    try:
        r = requests.get(ESPN_INJURIES, headers=S.BROWSER_HEADERS, timeout=30)
        r.raise_for_status()
        for team in r.json().get("injuries", []):
            for inj in team.get("injuries", []):
                name = inj.get("athlete", {}).get("displayName", "")
                rows.append({"key": name_key(name), "status": str(inj.get("status", "")),
                             "note": inj.get("shortComment", "") or ""})
        print(f"Injury report: {len(rows)} players listed (ESPN).")
    except Exception as e:
        print(f"Couldn't read the injury report ({e.__class__.__name__}); only manual overrides apply.")
    rep = pd.DataFrame(rows, columns=["key", "status", "note"])

    path = os.path.join(S.DATA_DIR, "injury_overrides.csv")
    if not os.path.exists(path):
        pd.DataFrame(columns=["player", "status"]).to_csv(path, index=False)
    ov = pd.read_csv(path)
    if len(ov):
        ov = pd.DataFrame({"key": ov.player.map(name_key), "status": ov.status.astype(str),
                           "note": "manual override"})
        rep = pd.concat([rep[~rep.key.isin(ov.key)], ov], ignore_index=True)
        print(f"Manual overrides: {len(ov)} players.")
    rep["out"] = rep.status.str.lower().str.startswith(OUT_STATUSES)
    return rep.drop_duplicates("key", keep="last").set_index("key")


def recency(dates, as_of, half_life):
    days = (pd.Timestamp(as_of) - pd.to_datetime(dates)).dt.days.clip(lower=0)
    return 0.5 ** (days.to_numpy() / half_life)


# ----------------------------------------------------------------- model
def per_minute_rates(players, prev1, prev2, log, as_of):
    """Each player's expected stats per minute, history blended with this season."""
    hist = pd.concat([prev1.assign(w=LAST_SEASON_WEIGHT), prev2.assign(w=TWO_SEASONS_WEIGHT)])
    for c in ["MIN"] + STATS:
        hist[c] = pd.to_numeric(hist[c], errors="coerce").fillna(0) * hist.w
    hist = hist.groupby("PLAYER_ID")[["MIN"] + STATS].sum()
    lg = prev1[STATS].sum() / max(prev1.MIN.sum(), 1)

    cur = log.assign(w=recency(log.GAME_DATE, as_of, RATE_HALF_LIFE)) if len(log) else log.assign(w=[])
    for c in ["MIN"] + STATS:
        cur[c] = cur[c].astype(float) * cur.w
    cur = cur.groupby("PLAYER_ID")[["MIN"] + STATS].sum()

    out = pd.DataFrame(index=players.PLAYER_ID)
    h = hist.reindex(out.index)
    c = cur.reindex(out.index).fillna(0)
    rookie = h.MIN.isna() | (h.MIN == 0)
    for s in STATS:
        base = np.where(rookie, ROOKIE_RATE * lg[s], lg[s])
        hist_rate = (h[s].fillna(0) + HIST_PRIOR_MIN[s] * base) / (h.MIN.fillna(0) + HIST_PRIOR_MIN[s])
        out[s] = (c[s] + CUR_PRIOR_MIN[s] * hist_rate) / (c.MIN + CUR_PRIOR_MIN[s])
    out["rookie"] = rookie.to_numpy()
    return out


def minutes_baseline(players, log, prev1, vb, as_of):
    """
    Minutes per game when he plays (baseline) and how often he plays (participation),
    each a preseason expectation blended with this season, weighted toward recent games.
    """
    p = players.copy()
    prior = p.key.map(vb) if vb is not None else pd.Series(np.nan, index=p.index)
    last = prev1.assign(mpg=prev1.MIN / prev1.GP.clip(lower=1)).groupby("PLAYER_ID").mpg.max()
    p["prior_mpg"] = prior.fillna(p.PLAYER_ID.map(last)).fillna(UNKNOWN_MPG)
    prior_part = np.where(p.prior_mpg >= 15, 1.0, 0.6)

    if len(log):
        lg_ = log.assign(w=recency(log.GAME_DATE, as_of, MIN_HALF_LIFE))
        on_team = lg_.merge(p[["PLAYER_ID", "TEAM_ID"]], on=["PLAYER_ID", "TEAM_ID"])
        played = on_team.groupby("PLAYER_ID").agg(w_games=("w", "sum"),
                                                  w_min=("MIN", lambda m: (m * on_team.loc[m.index, "w"]).sum()),
                                                  first=("GAME_DATE", "min"))
        team_games = lg_.drop_duplicates(["TEAM_ID", "GAME_ID"])[["TEAM_ID", "GAME_DATE", "w"]]
        g = p.merge(played, left_on="PLAYER_ID", right_index=True, how="left")
        # team games since he joined (all of them if he hasn't played yet)
        tg = []
        for tid, first in zip(g.TEAM_ID, g["first"]):
            t = team_games[team_games.TEAM_ID == tid]
            tg.append(t.w.sum() if pd.isna(first) else t[t.GAME_DATE >= first].w.sum())
        g["w_team_games"] = tg
        w_games, w_min = g.w_games.fillna(0).to_numpy(), g.w_min.fillna(0).to_numpy()
        w_team = g.w_team_games.to_numpy()
    else:
        w_games = w_min = w_team = np.zeros(len(p))
    p["baseline"] = (w_min + MIN_PRIOR_GAMES * p.prior_mpg) / (w_games + MIN_PRIOR_GAMES)
    p["part"] = ((w_games + PART_PRIOR_GAMES * prior_part) / (w_team + PART_PRIOR_GAMES)).clip(0, 1)
    return p


def allocate(team):
    """
    Fit one team's active players to 240 minutes; returns minutes if he plays.
    Short of 240, everyone scales up. Over 240, the even cut is blended with Version B's depth
    chart, so starters keep more of their minutes and the end of the bench gives up more.
    """
    base, part = team.baseline.to_numpy(), team.part.to_numpy()
    cap = np.minimum(np.maximum(REDIST_CAP, base + 1), MAX_MIN)

    def total(s):
        return (part * np.minimum(base * s, cap)).sum()

    lo, hi = 0.3, 3.0
    if total(hi) < TEAM_MINUTES:
        return np.minimum(base * hi, cap)
    for _ in range(50):
        mid = (lo + hi) / 2
        lo, hi = (mid, hi) if total(mid) < TEAM_MINUTES else (lo, mid)
    even = np.minimum(base * (lo + hi) / 2, cap)
    if total(1.0) <= TEAM_MINUTES:
        return even
    # Version B fills a season's minutes, so work in season units (x 82 games) and convert back
    expected = pd.Series(np.minimum(base, cap) * part * V.GAMES)
    depth = V.allocate_minutes(expected).to_numpy() / V.GAMES
    depth = np.divide(depth, part, out=np.zeros_like(depth), where=part > 0)
    return np.minimum((1 - DEPTH_WEIGHT) * even + DEPTH_WEIGHT * depth, cap)


def blowout(mins, part, margin):
    """Shift some starter minutes to the bench when the projected margin is large."""
    cut = BLOWOUT_PER_PT * max(0.0, abs(margin) - BLOWOUT_START)
    if cut == 0:
        return mins
    starters = mins >= STARTER_MIN
    freed = (mins[starters] * cut * part[starters]).sum()
    out = mins.copy()
    out[starters] *= 1 - cut
    bench = ~starters
    if bench.any():
        exp = out[bench] * part[bench]
        out[bench] += freed * exp / exp.sum() / part[bench].clip(min=0.05)
    return np.minimum(out, MAX_MIN)


def simulate(proj, rng):
    """Simulate every player's game; returns {stat: array (players x sims)}."""
    m = proj.MIN.to_numpy()[:, None]
    sd = 0.15 * m + 2
    mins = np.clip(m + sd * rng.standard_normal((len(proj), N_SIMS)), 0, 48)
    sims = {}
    for s in STATS:
        mu = np.maximum(proj[f"rate_{s}"].to_numpy()[:, None] * mins, 1e-6)
        k = DISPERSION[s]
        sims[s] = rng.negative_binomial(k, k / (k + mu)).astype(float)
    for c, parts in COMBOS.items():
        sims[c] = sum(sims[p] for p in parts)
    tens = sum((sims[s] >= 10).astype(int) for s in ["PTS", "REB", "AST", "STL", "BLK"])
    sims["DD"] = (tens >= 2).astype(float)
    return sims


# ------------------------------------------------------------------ main
def load_team_ratings():
    path = os.path.join(PR.OUT_DIR, "team_ratings.csv")
    lv_path = os.path.join(PR.OUT_DIR, "league_levels.csv")
    if not (os.path.exists(path) and os.path.exists(lv_path)):
        raise SystemExit("Can't find output_power/team_ratings.csv. Run nba_power_ratings.py first.")
    lv = pd.read_csv(lv_path)
    if lv.as_of.iloc[0] != dt.date.today().isoformat():
        print(f"Note: team ratings are from {lv.as_of.iloc[0]}; run nba_power_ratings.py for today's.")
    lg_o, lg_p = float(lv.league_off_rtg.iloc[0]), float(lv.league_pace.iloc[0])
    r = pd.read_csv(path, encoding="utf-8-sig").set_index("TEAM_NAME")
    r["off"], r["dfn"], r["pace_abs"] = r.off_rtg - lg_o, lg_o - r.def_rtg, r.pace
    r["pace"] = r.pace_abs - lg_p
    return r, lg_o, lg_p


def version_b_minutes():
    path = os.path.join("output_b", f"player_projections_{S.SEASON}.csv")
    if not os.path.exists(path):
        print("Version B player projections not found; using last season's minutes as the starting point.")
        return None, None
    vb = pd.read_csv(path, encoding="utf-8-sig")
    vb["key"] = vb.player.map(name_key)
    vb = vb.drop_duplicates("key")
    return vb.set_index("key").proj_mpg, vb.set_index("key").proj_bpm


def write_workbook(table, ranges, label):
    if not os.path.exists(PR.CALCULATOR_PATH):
        return
    from openpyxl import load_workbook
    try:
        wb = load_workbook(PR.CALCULATOR_PATH)
    except PermissionError:
        print(f"Couldn't open {PR.CALCULATOR_PATH} - close it in Excel and run again.")
        return
    for name, df in [("Player Projections", table), ("Player Ranges", ranges)]:
        if name in wb.sheetnames:
            ws = wb[name]
            ws.delete_rows(1, ws.max_row)
        else:
            ws = wb.create_sheet(name)
        ws.append(list(df.columns))
        for row in df.itertuples(index=False):
            ws.append([None if pd.isna(v) else (v.item() if hasattr(v, "item") else v) for v in row])
        ws.cell(1, len(df.columns) + 2).value = label
        ws.freeze_panes = "B2"
        ws.column_dimensions["A"].width = 24
    try:
        wb.save(PR.CALCULATOR_PATH)
        print(f"Updated {PR.CALCULATOR_PATH} (Player Projections and Player Ranges).")
    except PermissionError:
        print(f"Couldn't save {PR.CALCULATOR_PATH} - close it in Excel and run again.")


def main():
    day = pd.Timestamp(sys.argv[1]) if len(sys.argv) > 1 else pd.Timestamp(dt.date.today())
    day_str = day.date().isoformat()
    os.makedirs(OUT_DIR, exist_ok=True)
    rng = np.random.default_rng(SEED)

    teams = S.load_teams()
    names = teams.set_index("TEAM_ID").TEAM_NAME
    sched = S.fetch_schedule(set(teams.TEAM_ID))
    if sched is None:
        raise SystemExit("Couldn't get the schedule.")
    sched["game_id"] = sched.game_id.astype(str)
    sched = PR.back_to_back_adjustment(sched)
    games = sched[pd.to_datetime(sched.date).dt.date == day.date()]
    if games.empty:
        raise SystemExit(f"No games on {day_str}.")
    print(f"{len(games)} games on {day_str}.")

    ratings, lg_o, lg_p = load_team_ratings()
    vb_mpg, vb_bpm = version_b_minutes()
    start = int(S.SEASON[:4])
    prev1, prev2 = prev_season_totals(start - 1), prev_season_totals(start - 2)
    log = season_game_log()
    log = log[log.GAME_DATE < day]
    inj = injury_report()

    playing = set(games.home_id) | set(games.away_id)
    players = current_players()
    players = players[players.TEAM_ID.isin(playing)].copy()
    players["key"] = players.PLAYER_NAME.map(name_key)
    players = minutes_baseline(players, log, prev1, vb_mpg, day)
    rates = per_minute_rates(players, prev1, prev2, log, day)
    players = players.merge(rates, left_on="PLAYER_ID", right_index=True)
    players["status"] = players.key.map(inj.status).fillna("")
    players["out"] = players.key.map(inj.out).fillna(False).astype(bool)
    players["bpm"] = players.key.map(vb_bpm) if vb_bpm is not None else np.nan

    rows = []
    for g in games.itertuples():
        h, a = ratings.loc[names[g.home_id]], ratings.loc[names[g.away_id]]
        m = PR.matchup(h, a, lg_o, lg_p, g.rest_adj)
        # out players' value (Version B rating above a replacement player, times his usual role)
        impact = {}
        for tid in (g.home_id, g.away_id):
            o = players[(players.TEAM_ID == tid) & players.out]
            impact[tid] = ((o.bpm.fillna(REPLACEMENT_BPM) - REPLACEMENT_BPM).clip(lower=0)
                           * o.prior_mpg * o.part / 48).sum()
        home_pts = m["home_pts"] - impact[g.home_id] / 2 + impact[g.away_id] / 2
        away_pts = m["away_pts"] - impact[g.away_id] / 2 + impact[g.home_id] / 2
        margin = home_pts - away_pts
        for tid, opp, side, pts, team_r in [(g.home_id, g.away_id, "Home", home_pts, h),
                                            (g.away_id, g.home_id, "Away", away_pts, a)]:
            t = players[(players.TEAM_ID == tid) & ~players.out].copy()
            if t.empty:
                continue
            mins = allocate(t)
            mins = blowout(mins, t.part.to_numpy(), margin)
            t["MIN"] = mins
            exp_min = mins * t.part.to_numpy()
            pace_f = m["poss"] / team_r.pace_abs
            raw_pts = (t.PTS * exp_min).sum()
            score_f = float(np.clip(pts / raw_pts, 0.8, 1.25)) if raw_pts > 0 else 1.0
            for s in STATS:
                t[f"rate_{s}"] = t[s] * (score_f if s in SCORING else pace_f)
            t["Team"], t["Opp"], t["Home/Away"] = ABBR.get(tid, tid), ABBR.get(opp, opp), side
            t["team_proj_pts"] = pts
            rows.append(t)
    proj = pd.concat(rows, ignore_index=True)
    proj = proj[(proj.MIN >= MIN_TO_SHOW) & ((proj.part >= 0.25) | (proj.status != ""))].reset_index(drop=True)

    sims = simulate(proj, rng)
    table = proj[["PLAYER_NAME", "Team", "Opp", "Home/Away", "status"]].rename(
        columns={"PLAYER_NAME": "Player", "status": "Status"})
    table["MIN"] = proj.MIN.round(1)
    for s in STATS + list(COMBOS):
        table[LABELS.get(s, s)] = sims[s].mean(axis=1).round(1)
    table["DD%"] = (sims["DD"].mean(axis=1) * 100).round(0)
    table = table.sort_values(["Team", "MIN"], ascending=[True, False])

    rng_rows = []
    for s in STATS + list(COMBOS):
        q = np.percentile(sims[s], [10, 25, 50, 75, 90], axis=1)
        rng_rows.append(pd.DataFrame({"Player": proj.PLAYER_NAME, "Team": proj.Team, "Opp": proj.Opp,
                                      "Stat": LABELS.get(s, s), "Mean": sims[s].mean(axis=1).round(1),
                                      "P10": q[0], "P25": q[1], "Median": q[2], "P75": q[3], "P90": q[4]}))
    ranges = pd.concat(rng_rows, ignore_index=True)

    table.to_csv(os.path.join(OUT_DIR, f"player_projections_{day_str}.csv"), index=False, encoding="utf-8-sig")
    ranges.to_csv(os.path.join(OUT_DIR, f"player_ranges_{day_str}.csv"), index=False, encoding="utf-8-sig")
    hist_path = os.path.join(OUT_DIR, "projection_history.csv")
    snap = table.assign(date=day_str, PLAYER_ID=proj.loc[table.index, "PLAYER_ID"].to_numpy())
    if os.path.exists(hist_path):
        old = pd.read_csv(hist_path, encoding="utf-8-sig")
        snap = pd.concat([old[old.date != day_str], snap], ignore_index=True)
    snap.to_csv(hist_path, index=False, encoding="utf-8-sig")
    # ranges too, so nba_grade.py can check whether the ranges are the right width
    rh_path = os.path.join(OUT_DIR, "range_history.csv")
    rsnap = ranges.assign(date=day_str, PLAYER_ID=np.tile(proj.PLAYER_ID.to_numpy(), len(STATS) + len(COMBOS)))
    if os.path.exists(rh_path):
        old = pd.read_csv(rh_path, encoding="utf-8-sig")
        rsnap = pd.concat([old[old.date != day_str], rsnap], ignore_index=True)
    rsnap.to_csv(rh_path, index=False, encoding="utf-8-sig")

    out_list = players[players.out][["PLAYER_NAME", "TEAM_ID", "status"]]
    if len(out_list):
        print("\nNot projected (Out/Doubtful):")
        for r in out_list.itertuples():
            print(f"  {r.PLAYER_NAME} ({ABBR.get(r.TEAM_ID, r.TEAM_ID)}) - {r.status}")
    label = f"Projections for {day_str}"
    print(f"\n{label} (minutes and stats assume the player plays)\n")
    print(table.to_string(index=False))
    write_workbook(table, ranges, label)
    print(f"\nSaved results to ./{OUT_DIR}")


if __name__ == "__main__":
    main()
